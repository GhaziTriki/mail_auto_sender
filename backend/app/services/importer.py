"""Contacts file import (spec 8.1)."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import shutil
from pathlib import Path

from ..config import get_settings
from ..errors import Unprocessable


def run_dir(run_id) -> Path:
    p = Path(get_settings().data_dir) / "runs" / str(run_id)
    p.mkdir(parents=True, exist_ok=True)
    return p


def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def detect_encoding(raw: bytes) -> str:
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    try:
        raw.decode("utf-8")
        return "utf-8-sig"
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(raw).best()
        if best is not None and best.encoding and best.encoding.lower() not in ("ascii",):
            enc = best.encoding.lower()
            if enc in ("utf_8", "utf-8"):
                return "utf-8-sig"
            try:
                raw.decode(enc)
                return enc
            except Exception:
                print()
    except Exception:
        print()
    return "latin-1"


def detect_delimiter(text: str) -> str:
    sample = text[: 64 * 1024]
    try:
        d = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        return d
    except csv.Error:
        return ","


def normalize_headers(headers: list[str]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, h in enumerate(headers, start=1):
        name = (h or "").strip()
        if not name:
            name = f"column_{i}"
        if name in seen:
            seen[name] += 1
            candidate = f"{name}_{seen[name]}"
            while candidate in seen:
                seen[name] += 1
                candidate = f"{name}_{seen[name]}"
            seen[candidate] = 1
            out.append(candidate)
        else:
            seen[name] = 1
            out.append(name)
    return out


def _parse_csv(
    raw: bytes, encoding: str | None, delimiter: str | None
) -> tuple[list[str], list[list[str]], str, str]:
    enc = encoding or detect_encoding(raw)
    try:
        text = raw.decode(enc)
    except (UnicodeDecodeError, LookupError) as e:
        raise Unprocessable(f"Cannot decode file with encoding {enc}: {e}", code="bad_encoding") from e
    delim = delimiter or detect_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    rows = list(reader)
    if not rows:
        raise Unprocessable("The file is empty", code="empty_file")
    return rows[0], rows[1:], enc, delim


def sheet_names(raw: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _parse_xlsx(raw: bytes, sheet: str | None) -> tuple[list[str], list[list[str]], str | None, list[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise Unprocessable(f"Cannot read the XLSX file: {e}", code="bad_xlsx") from e
    try:
        names = list(wb.sheetnames)
        if not names:
            raise Unprocessable("The workbook has no sheets", code="empty_file")
        chosen = sheet if sheet in names else names[0]
        ws = wb[chosen]
        rows = [[_cell_str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if not rows:
        raise Unprocessable("The sheet is empty", code="empty_file")
    return rows[0], rows[1:], chosen, names


def parse_source(
    raw: bytes,
    fmt: str,
    *,
    encoding: str | None = None,
    delimiter: str | None = None,
    sheet: str | None = None,
) -> dict:
    """Parse and normalise. Returns dict(columns, rows(list of dict), encoding, delimiter, sheet, sheets)."""
    st = get_settings()
    if len(raw) > st.max_source_mb * 1024 * 1024:
        raise Unprocessable(f"File is larger than {st.max_source_mb} MB", code="file_too_large")
    sheets: list[str] = []
    if fmt == "csv":
        header, body, enc, delim = _parse_csv(raw, encoding, delimiter)
        used_sheet = None
    elif fmt == "xlsx":
        header, body, used_sheet, sheets = _parse_xlsx(raw, sheet)
        enc, delim = None, None
    else:
        raise Unprocessable("Only .csv and .xlsx files are accepted", code="bad_format")
    columns = normalize_headers([h.strip() if isinstance(h, str) else str(h) for h in header])
    rows: list[dict] = []
    for r in body:
        vals = [(c.strip() if isinstance(c, str) else _cell_str(c)) for c in r]
        if not any(vals):
            continue
        vals = (vals + [""] * len(columns))[: len(columns)]
        rows.append(dict(zip(columns, vals, strict=True)))
    if len(rows) > st.max_rows:
        raise Unprocessable(f"File has more than {st.max_rows} rows", code="too_many_rows")
    return {
        "columns": columns,
        "rows": rows,
        "encoding": enc,
        "delimiter": delim,
        "sheet": used_sheet,
        "sheets": sheets,
    }


def detect_format(filename: str) -> str:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext == ".csv":
        return "csv"
    if ext == ".xlsx":
        return "xlsx"
    raise Unprocessable("Only .csv and .xlsx files are accepted", code="bad_format")


def save_source(run_id, filename: str, raw: bytes, *, encoding=None, delimiter=None, sheet=None) -> dict:
    """Parse, copy original into the run dir and write parsed.jsonl. Returns the run field updates."""
    fmt = detect_format(filename)
    parsed = parse_source(raw, fmt, encoding=encoding, delimiter=delimiter, sheet=sheet)
    d = run_dir(run_id)
    for old in d.glob("source.*"):
        old.unlink()
    src_path = d / f"source.{fmt}"
    src_path.write_bytes(raw)
    write_parsed(run_id, parsed["rows"])
    return {
        "source_file": {
            "path": str(src_path),
            "original_name": os.path.basename(filename),
            "sha256": sha256_bytes(raw),
            "size": len(raw),
            "format": fmt,
            "delimiter": parsed["delimiter"],
            "encoding": parsed["encoding"],
            "sheet": parsed["sheet"],
        },
        "columns": parsed["columns"],
        "row_count": len(parsed["rows"]),
        "_preview": {"rows": parsed["rows"][:20], "sheets": parsed["sheets"]},
    }


def write_parsed(run_id, rows: list[dict]) -> None:
    p = run_dir(run_id) / "parsed.jsonl"
    tmp = p.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        for i, r in enumerate(rows):
            f.write(json.dumps({"_i": i, **r}, ensure_ascii=False) + "\n")
    os.replace(tmp, p)


def read_parsed(run_id):
    p = run_dir(run_id) / "parsed.jsonl"
    if not p.exists():
        return
    with open(p, encoding="utf-8") as f:
        for line in f:
            if line.strip():
                yield json.loads(line)


def save_cv(run_id, filename: str, raw: bytes) -> dict:
    st = get_settings()
    if not raw.startswith(b"%PDF"):
        raise Unprocessable("The CV must be a PDF file", code="cv_not_pdf")
    if len(raw) > st.max_cv_mb * 1024 * 1024:
        raise Unprocessable(f"The CV must be at most {st.max_cv_mb} MB", code="cv_too_large")
    path = run_dir(run_id) / "cv.pdf"
    path.write_bytes(raw)
    return {
        "path": str(path),
        "original_name": os.path.basename(filename) or "cv.pdf",
        "sha256": sha256_bytes(raw),
        "size": len(raw),
    }


def delete_run_files(run_id) -> None:
    p = Path(get_settings().data_dir) / "runs" / str(run_id)
    shutil.rmtree(p, ignore_errors=True)
